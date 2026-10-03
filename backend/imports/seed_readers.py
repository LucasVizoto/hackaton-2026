"""Read and reconcile a complete private package before any database writes."""

import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from django.core.exceptions import ValidationError as DjangoValidationError
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import ValidationError

from catalog.models import Depot, Equipment, Product, ProductDeposit, Supplier, Worker
from labor.constants import RATE_TABLE
from labor.models import BulletinLine, DailyBulletin
from receiving.xml_parser import parse_invoice_xml

from .models import HistoricalLaborDay, HistoricalMovement, HistoricalStock, HistoricalWorkerDay, SourceFile
from .services import (
    READERS, SOURCE_FILES, PrivateDataError, decimal_value, json_value, normalized,
    record_outcomes, text,
)

SEED_VERSION = "hackathon-2026-v1"
STOCK_FILES = {
    "02_cadastros/estoque_por_armazem/estoque_defensivos.xlsx": "MATDefe",
    "02_cadastros/estoque_por_armazem/estoque_fertilizantes.xlsx": "MATFerti",
    "02_cadastros/estoque_por_armazem/estoque_geral.xlsx": "MATGeral",
    "02_cadastros/estoque_por_armazem/estoque_maquinas.xlsx": "MATMaq",
}
PAYROLL_FILES = (
    "04_mao_de_obra/chapas_por_dia_2025.xlsx",
    "04_mao_de_obra/chapas_por_dia_2026.xlsx",
)
EQUIPMENT_FILE = "05_operacao/equipamentos_descarga.xlsx"
BULLETIN_FILE = SOURCE_FILES["workers"]
REFERENCE_FILES = (
    "LEIA-ME.md", "05_operacao/registro_manual_recebimento.pdf",
    "05_operacao/formulario_registro_em_branco.doc",
    "06_especificacao/recebimento_inteligente_cocapec.pdf",
    "06_especificacao/mapa_relacoes_sap.jpeg",
)
MONTHS = {name: index for index, name in enumerate(
    ("janeiro", "fevereiro", "marco", "abril", "maio", "junho", "julho",
     "agosto", "setembro", "outubro", "novembro", "dezembro"), 1
)}


def file_digest(path):
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def source_rows(path):
    """Preserve column addresses, formulas and cached values, including non-data rows."""
    formulas = load_workbook(path, read_only=True, data_only=False)
    cached = load_workbook(path, read_only=True, data_only=True)
    result = []
    try:
        for sheet in formulas:
            values = cached[sheet.title].iter_rows(values_only=True)
            for row_number, (row, cache) in enumerate(zip(sheet.iter_rows(), values, strict=True), 1):
                cells = {}
                for column, (cell, value) in enumerate(zip(row, cache, strict=True), 1):
                    if cell.value is not None:
                        cells[get_column_letter(column)] = {
                            "value": json_value(cell.value), "cached": json_value(value),
                            "is_formula": cell.data_type == "f",
                        }
                if cells:
                    result.append({"sheet": sheet.title, "row": row_number, "original": cells})
    finally:
        formulas.close()
        cached.close()
    return result


def read_stock(path, depot):
    book = load_workbook(path, read_only=True, data_only=True)
    result = []
    try:
        sheet = book.active
        rows = sheet.iter_rows(values_only=True)
        next(rows)
        header = [normalized(v) for v in next(rows)]
        if (len(header) < 3 or header[0] not in {"n.item", "codigo do produto"}
                or header[1] != "descricao"
                or header[2] not in {"qtd.sap", "quantidade em estoque"}):
            raise PrivateDataError("Colunas de estoque incompatíveis.")
        for index, row in enumerate(rows, 3):
            if not row[0]:
                continue  # Titles/groups remain in SourceRow and in the original file.
            problems = []
            quantity = decimal_value(row[2], problems, "quantity")
            if quantity is not None and (abs(quantity) >= Decimal("1e18") or quantity.as_tuple().exponent < -6):
                problems.append("quantity_outside_precision")
                quantity = None
            result.append({
                "source_sheet": sheet.title, "source_row": index,
                "depot_code": depot, "product_code": text(row[0]),
                "description": text(row[1]), "quantity": quantity,
                "original": {get_column_letter(j + 1): json_value(v) for j, v in enumerate(row)},
                "problems": problems,
            })
    finally:
        book.close()
    if not result:
        raise PrivateDataError("Nenhuma posição de estoque encontrada.")
    return result


def read_equipment(path):
    book = load_workbook(path, read_only=True, data_only=True)
    records = []
    codes = {
        "paleteira manual": "PALETEIRA_MANUAL",
        "empilhadeira a gas": "EMPILHADEIRA_GAS",
        "carrinho manual": "CARRINHO_MANUAL",
        "transpaleteira eletrica": "TRANSPALETEIRA_ELETRICA",
    }
    try:
        rows = list(book.active.iter_rows(values_only=True))
        if len(rows) < 3 or [normalized(v) for v in rows[2][:2]] != ["equipamento", "utilizacao"]:
            raise PrivateDataError("Colunas de equipamentos incompatíveis.")
        for index, row in enumerate(rows[3:], 4):
            if not row[0]:
                continue
            code = codes.get(normalized(row[0]))
            if not code:
                raise PrivateDataError("Tipo de equipamento não reconhecido; atualize o leitor.")
            records.append({"code": code, "name": text(row[0]), "purpose": text(row[1]),
                            "warehouse_id": None, "mobile": None, "source_row": index})
    finally:
        book.close()
    if not records or len({r["code"] for r in records}) != len(records):
        raise PrivateDataError("Catálogo de equipamentos vazio ou repetido.")
    return records


def read_bulletin(path, workers):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.active
        declared_day = sheet.cell(5, 9).value
        if not isinstance(declared_day, (datetime, date)):
            raise PrivateDataError("Data do boletim preenchido não encontrada.")
        if "adubo" not in normalized(sheet.cell(3, 2).value):
            raise PrivateDataError("Local do boletim preenchido não reconhecido.")
        lines, participants = [], []
        for index, (code, label, _) in enumerate(RATE_TABLE):
            row = 10 + index * 2
            declared_category = normalized(sheet.cell(row, 1).value).replace(" ", "")
            allowed = {normalized(label).replace(" ", "")}
            if code == "ACESSORIOS":
                allowed.add("acessoriosagropec")
            if declared_category not in allowed:
                raise PrivateDataError("Categorias ou linhas do boletim incompatíveis com o leitor.")
            problems = []
            price = decimal_value(sheet.cell(row, 8).value, problems, "price")
            quantities = {}
            for key, column in (("unloading", 3), ("removal", 4), ("transfer", 6)):
                value = sheet.cell(row, column).value
                quantities[key] = Decimal(0) if value is None else decimal_value(value, problems, key)
            if problems or price < 0 or price.as_tuple().exponent < -4 or any(
                v < 0 or v.as_tuple().exponent < -4 for v in quantities.values()
            ):
                raise PrivateDataError("Tarifa ou quantidade inválida no boletim.")
            lines.append({"category": code, "label": label, "price": price, **quantities})
        known = {r["key"] for r in workers}
        for row in range(47, 57):
            for column, half in ((3, 4), (8, 9)):
                registration = text(sheet.cell(row, column).value)
                if not registration:
                    continue
                marker = normalized(sheet.cell(row, half).value)
                if registration not in known or marker not in {"", "meia"}:
                    raise PrivateDataError("Participante do boletim sem matrícula ou fração inequívoca.")
                participants.append({"registration": registration,
                                     "fraction": Decimal("0.5") if marker else Decimal(1)})
        if not participants or len({p["registration"] for p in participants}) != len(participants):
            raise PrivateDataError("Equipe do boletim vazia ou repetida.")
        problems = []
        floor = decimal_value(sheet.cell(38, 8).value, problems, "floor")
        if problems or floor < 0 or floor.as_tuple().exponent < -4:
            raise PrivateDataError("Piso do boletim inválido.")
        return {"day": declared_day.date() if isinstance(declared_day, datetime) else declared_day,
                "warehouse_code": "ADUBO", "floor": floor, "lines": lines,
                "participants": participants, "declared_half_floor": json_value(sheet.cell(40, 8).value)}
    finally:
        book.close()


def title_matches_day(title, day):
    title = normalized(title)
    month = next((number for name, number in MONTHS.items() if name in title), None)
    year = re.search(r"\b20\d{2}\b", title)
    return month == day.month and year is not None and int(year.group()) == day.year


def read_payroll(path):
    book = load_workbook(path, read_only=True, data_only=True)
    records, header_count = [], 0
    try:
        for sheet in book:
            columns = {}
            for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if len(row) > 1 and normalized(row[1]) == "nome":
                    header_count += 1
                    columns = {j: v for j, v in enumerate(row) if j >= 12 and (
                        isinstance(v, (date, datetime)) or isinstance(v, (int, float)) and 1 <= v <= 31
                    )}
                    if not columns:
                        raise PrivateDataError("Cabeçalho da folha sem colunas de dias reconhecíveis.")
                    continue
                identifier = re.match(r"^(CHAPA_\d+)(?:\s|$)", text(row[1])) if len(row) > 1 else None
                if not identifier:
                    continue
                if not columns:
                    raise PrivateDataError("Pessoa na folha sem cabeçalho de datas.")
                for column, declared in columns.items():
                    problems = []
                    day = declared.date() if isinstance(declared, datetime) else declared
                    if not isinstance(day, date) or day.year not in {2025, 2026}:
                        day = None
                        problems.append("invalid_declared_date")
                    elif not title_matches_day(sheet.title, day):
                        problems.append("sheet_title_date_mismatch")
                    raw = row[column] if column < len(row) else None
                    amount = decimal_value(raw, problems, "payroll_paid")
                    if amount is not None and (abs(amount) >= Decimal("1e14") or amount.as_tuple().exponent < -6):
                        amount = None
                        problems.append("payroll_outside_precision")
                    if amount is not None and amount < 0:
                        problems.append("negative_payroll_preserved")
                    records.append({
                        "source_sheet": sheet.title, "source_row": index,
                        "source_column": column + 1, "source_identifier": identifier.group(1),
                        "declared_date": str(json_value(declared)), "day": day,
                        "location": text(row[0]), "payroll_paid": amount,
                        "usable": day is not None and amount is not None and amount >= 0,
                        "original": {"name": json_value(row[1]), "date": json_value(declared),
                                     "value": json_value(raw), "cell": f"{get_column_letter(column + 1)}{index}"},
                        "problems": problems,
                    })
    finally:
        book.close()
    if not header_count or not records:
        raise PrivateDataError("Folha individual sem registros reconhecíveis.")
    return records


def reconcile_payroll(payroll):
    groups = defaultdict(list)
    for records in payroll.values():
        for record in records:
            if record["day"] is not None:
                groups[(record["source_identifier"], record["day"])].append(record)
    for records in groups.values():
        if len(records) < 2:
            continue
        if len({(r["payroll_paid"], r["location"]) for r in records}) > 1:
            for record in records:
                record["usable"] = False
                record["problems"].append("conflicting_worker_day")
        else:
            ordered = sorted(records, key=lambda r: (
                not title_matches_day(r["source_sheet"], r["day"]),
                r["source_sheet"], r["source_row"], r["source_column"],
            ))
            for record in ordered[1:]:
                record["usable"] = False
                record["problems"].append("duplicate_worker_day_preserved")


def validate_values(model, attrs):
    """Validate persisted scalar fields without querying an unmigrated database."""
    instance = model(**attrs)
    for field_name, value in attrs.items():
        field = model._meta.get_field(field_name)
        if field.is_relation or field.get_internal_type() == "JSONField":
            continue
        try:
            field.clean(value, instance)
        except DjangoValidationError:
            raise PrivateDataError(f"Valor incompatível com o esquema: {model._meta.label}.{field_name}.") from None


def validate_catalog(records, model):
    known = {}
    for record in records:
        attrs = record["attrs"]
        key_field = "registration" if model is Worker else "code"
        key = attrs[key_field]
        if not key:
            continue
        if key in known and known[key] != attrs:
            raise PrivateDataError("Atributos conflitantes no catálogo; corrija a fonte antes do seed.")
        known[key] = attrs
        validate_values(model, attrs)


def prepare_seed(directory):
    directory = Path(directory).expanduser().resolve()
    if not directory.is_dir():
        raise PrivateDataError("Pasta privada não encontrada. Configure PRIVATE_DATA_DIR ou --path.")
    required = set(SOURCE_FILES.values()) | set(STOCK_FILES) | set(PAYROLL_FILES) | {EQUIPMENT_FILE} | set(REFERENCE_FILES)
    missing = sorted(p for p in required if not (directory / p).is_file())
    if missing:
        raise PrivateDataError("Fontes obrigatórias ausentes: " + ", ".join(missing))
    xml_paths = sorted(directory.glob("01_notas_fiscais/xml/*.xml"))
    pdf_paths = sorted(directory.glob("01_notas_fiscais/danfe_pdf/*.pdf"))
    if not xml_paths or not pdf_paths:
        raise PrivateDataError("Pastas de XML e DANFE devem conter documentos.")
    manifest = []
    legacy_paths = {path: kind for kind, path in SOURCE_FILES.items()}
    for path in sorted(directory.rglob("*")):
        if not path.is_file():
            continue
        if not path.resolve().is_relative_to(directory):
            raise PrivateDataError("Fonte fora da pasta privada configurada.")
        relative = path.relative_to(directory).as_posix()
        validate_values(SourceFile, {"relative_path": relative})
        kind = legacy_paths.get(relative, "reference")
        if relative in STOCK_FILES:
            kind = "stock"
        elif relative in PAYROLL_FILES:
            kind = "worker_days"
        elif relative == EQUIPMENT_FILE:
            kind = "equipment"
        elif path in xml_paths:
            kind = "invoice_xml"
        elif path in pdf_paths:
            kind = "invoice_pdf"
        if path.suffix.lower() == ".pdf":
            with path.open("rb") as stream:
                if not stream.read(5).startswith(b"%PDF-"):
                    raise PrivateDataError("Arquivo PDF inválido na pasta privada.")
        manifest.append({"path": relative, "sha256": file_digest(path),
                         "size": path.stat().st_size, "kind": kind})
    manifests = {item["path"]: item for item in manifest}
    legacy = {}
    for kind, relative in SOURCE_FILES.items():
        records, summary, issues = READERS[kind](directory / relative)
        if not records:
            raise PrivateDataError(f"Fonte obrigatória sem registros: {relative}.")
        legacy[kind] = {"records": records, "summary": summary, "issues": issues,
                        "digest": manifests[relative]["sha256"]}
    for kind, model in (("products", Product), ("suppliers", Supplier), ("workers", Worker)):
        validate_catalog(legacy[kind]["records"], model)
    for kind, model in (("movements", HistoricalMovement), ("labor_days", HistoricalLaborDay)):
        for record in legacy[kind]["records"]:
            validate_values(model, record)
    for record in legacy["products"]["records"]:
        if record["depot"]:
            validate_values(ProductDeposit, {"depot": record["depot"]})
    stocks = {p: read_stock(directory / p, depot) for p, depot in STOCK_FILES.items()}
    product_codes = {r["attrs"]["code"] for r in legacy["products"]["records"]}
    for records in stocks.values():
        for record in records:
            if record["product_code"] not in product_codes:
                record["problems"].append("product_missing_from_current_catalog")
    payroll = {p: read_payroll(directory / p) for p in PAYROLL_FILES}
    reconcile_payroll(payroll)
    equipment = read_equipment(directory / EQUIPMENT_FILE)
    bulletin = read_bulletin(directory / BULLETIN_FILE, legacy["workers"]["records"])
    for records in stocks.values():
        for record in records:
            validate_values(HistoricalStock, {k: v for k, v in record.items() if k != "depot_code"})
    for records in payroll.values():
        for record in records:
            validate_values(HistoricalWorkerDay, record)
    for record in equipment:
        validate_values(Equipment, {k: v for k, v in record.items() if k not in {"source_row", "warehouse_id"}})
    for line in bulletin["lines"]:
        validate_values(BulletinLine, {k: v for k, v in line.items() if k != "label"})
    validate_values(DailyBulletin, {"reference_date": bulletin["day"], "floor_per_day": bulletin["floor"]})
    documents, movements = defaultdict(set), defaultdict(set)
    for record in legacy["suppliers"]["records"]:
        documents[re.sub(r"\D", "", record["attrs"]["document"])].add(record["key"])
    for record in legacy["movements"]["records"]:
        movements[record["invoice_key"]].add(record["supplier_code"])
    invoices = defaultdict(list)
    for path in xml_paths:
        try:
            extracted = parse_invoice_xml(path.read_bytes())
        except ValidationError:
            raise PrivateDataError("XML estruturalmente inválido; confira os documentos privados.") from None
        if not extracted["access_key"]:
            raise PrivateDataError("XML do conjunto histórico sem chave de acesso.")
        invoices[extracted["access_key"]].append({"path": path.relative_to(directory).as_posix(), "extracted": extracted})
    invoice_records, file_problems = [], {}
    for key, versions in sorted(invoices.items()):
        versions.sort(key=lambda v: (Path(v["path"]).stem != key, v["path"]))
        canonical = versions[0]
        candidates = documents[re.sub(r"\D", "", canonical["extracted"]["issuer"]["document"])]
        if movements[key]:
            candidates = candidates & movements[key]
        problems = []
        if len(candidates) != 1:
            problems.append("supplier_document_movement_conflict")
        if any(v["extracted"] != canonical["extracted"] for v in versions):
            problems.append("conflicting_invoice_versions")
        for index, version in enumerate(versions):
            file_problems[version["path"]] = [*problems, *(["duplicate_invoice_source_preserved"] if index else [])]
        invoice_records.append({"key": key, "supplier_code": next(iter(candidates)) if len(candidates) == 1 else None,
                                "primary": canonical["path"], "extracted": canonical["extracted"],
                                "paths": [v["path"] for v in versions], "problems": problems})
    by_key = {r["key"]: r for r in invoice_records}
    supplier_codes = {r["key"] for r in legacy["suppliers"]["records"] if r["key"]}
    for path in pdf_paths:
        relative = path.relative_to(directory).as_posix()
        match = re.fullmatch(r"(\d{44})(?: \(\d+\))?", path.stem)
        if not match:
            file_problems[relative] = ["unrecognized_pdf_access_key"]
            continue
        key = match.group(1)
        if key not in by_key:
            candidates = movements[key]
            problems = [] if len(candidates) == 1 and candidates <= supplier_codes else ["pdf_supplier_unresolved"]
            record = {"key": key, "supplier_code": next(iter(candidates)) if len(candidates) == 1 else None,
                      "primary": relative, "extracted": {}, "paths": [], "problems": problems}
            invoice_records.append(record)
            by_key[key] = record
        by_key[key]["paths"].append(relative)
        file_problems[relative] = list(by_key[key]["problems"])
    raw_rows = {p: source_rows(directory / p) for p in (*STOCK_FILES, *PAYROLL_FILES, EQUIPMENT_FILE, BULLETIN_FILE)}
    depots = {r["depot"] for r in legacy["products"]["records"] if r["depot"]}
    depots.update(r["depot"] for r in legacy["movements"]["records"] if r["depot"])
    depots.update(STOCK_FILES.values())
    for code in depots:
        validate_values(Depot, {"code": code})
    report = {"version": SEED_VERSION, "status": "validated", "files": len(manifest),
              "files_by_kind": dict(Counter(m["kind"] for m in manifest)),
              "datasets": {kind: {"row_count": len(p["records"]), "summary": {**p["summary"], **record_outcomes(p["records"])},
                                   "issues": p["issues"]} for kind, p in legacy.items()},
              "stock_rows": sum(map(len, stocks.values())), "stock_missing_products": sum(
                  "product_missing_from_current_catalog" in r["problems"] for rows in stocks.values() for r in rows),
              "worker_day_cells": sum(map(len, payroll.values())),
              "worker_day_usable_cells": sum(r["usable"] for rows in payroll.values() for r in rows),
              "worker_day_issues": dict(Counter(p for rows in payroll.values() for r in rows for p in r["problems"])),
              "invoice_keys": len(invoice_records), "invoices_ready": sum(not r["problems"] for r in invoice_records),
              "invoices_pending": sum(bool(r["problems"]) for r in invoice_records),
              "depots": len(depots), "equipment": len(equipment), "bulletins": 1}
    return {"directory": directory, "manifest": manifest,
            "manifest_hash": hashlib.sha256(json.dumps(manifest, sort_keys=True).encode()).hexdigest(),
            "legacy": legacy, "stocks": stocks, "payroll": payroll, "equipment": equipment,
            "bulletin": bulletin, "invoices": invoice_records, "file_problems": file_problems,
            "raw_rows": raw_rows, "depots": depots, "report": report}
