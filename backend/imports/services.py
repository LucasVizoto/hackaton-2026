"""Small, versioned readers. Source bytes stay in a configured private directory."""

import csv
import hashlib
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.db import connection, transaction
from openpyxl import load_workbook

from catalog.models import Product, ProductDeposit, Supplier, Worker

from .models import HistoricalLaborDay, HistoricalMovement, ImportBatch, SourceRow

IMPORTER_VERSION = "1.1"
SOURCE_FILES = {
    "products": "02_cadastros/produtos.xlsx",
    "suppliers": "02_cadastros/fornecedores.xlsx",
    "workers": "05_operacao/boletim_diario_chapas.xlsx",
    "movements": "03_movimentacao/pedido_recebimento_notafiscal.xlsx",
    "labor_days": "04_mao_de_obra/chapas_por_dia.csv",
}


class PrivateDataError(ValueError):
    pass


def text(value):
    if value is None:
        return ""
    if isinstance(value, (int, float, Decimal)) and Decimal(str(value)) == Decimal(str(value)).to_integral():
        return str(int(value))
    return str(value).strip()


def json_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def original_values(headers, values):
    return {headers[i]: json_value(value) for i, value in enumerate(values)}


def normalized(value):
    return "".join(c for c in unicodedata.normalize("NFKD", text(value).lower()) if not unicodedata.combining(c)).strip()


def decimal_value(value, problems, name):
    if value is None or text(value) == "":
        problems.append(f"missing_{name}")
        return None
    try:
        result = Decimal(str(value).replace(",", "."))
        if not result.is_finite() or abs(result) >= Decimal("1e18"):
            raise InvalidOperation
        return result
    except (InvalidOperation, ValueError):
        problems.append(f"invalid_{name}")
        return None


def date_value(value, problems, name):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(text(value)[:10])
    except ValueError:
        problems.append(f"missing_or_invalid_{name}")
        return None


def workbook_rows(path, expected_headers):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.active
        rows = sheet.iter_rows(values_only=True)
        headers = [text(v) for v in next(rows)]
        required = [normalized(h) for h in expected_headers]
        actual = [normalized(h) for h in headers]
        if actual[: len(required)] != required:
            raise PrivateDataError("Colunas incompatíveis com a versão do leitor; nenhuma linha foi gravada.")
        for index, values in enumerate(rows, 2):
            if any(value is not None for value in values):
                yield sheet.title, index, headers, values
    finally:
        book.close()


def read_products(path):
    records, issues, attributes = [], Counter(), {}
    headers = ["Nº do item", "Descrição do item", "Unidade de medida", "Peso", "Nome do grupo", "Grupo", "Descrição", "depósito"]
    for sheet, row, heads, values in workbook_rows(path, headers):
        code, depot = text(values[0]), text(values[7])
        problems = []
        if not code:
            problems.append("missing_product_code")
        if not depot:
            problems.append("missing_depot")
        attrs = {"code": code, "name": text(values[1]), "unit": text(values[2]), "weight": decimal_value(values[3], problems, "weight"), "group": text(values[5])}
        if code in attributes and attributes[code] != attrs:
            problems.append("conflicting_product_attributes")
        attributes.setdefault(code, attrs)
        issues.update(problems)
        records.append({"sheet": sheet, "row": row, "key": f"{code}|{depot}", "original": original_values(heads, values), "attrs": attrs, "depot": depot, "problems": problems})
    summary = {"unique_products": len({r["attrs"]["code"] for r in records if r["attrs"]["code"]}), "unique_product_depot_pairs": len({r["key"] for r in records})}
    if summary["unique_product_depot_pairs"] != len(records):
        issues["duplicate_product_depot_rows"] = len(records) - summary["unique_product_depot_pairs"]
        seen = set()
        for record in records:
            if record["key"] in seen:
                record["problems"].append("duplicate_product_depot_row_preserved")
            seen.add(record["key"])
    return records, summary, dict(issues)


def read_suppliers(path):
    records, issues, documents = [], Counter(), defaultdict(set)
    for sheet, row, heads, values in workbook_rows(path, ["COD", "FORNECEDOR", "CNPJ"]):
        code, document = text(values[0]), text(values[2])
        problems = []
        if not code:
            problems.append("missing_supplier_code")
        if not document:
            problems.append("missing_supplier_document")
        issues.update(problems)
        if document:
            documents[document].add(code)
        records.append({"sheet": sheet, "row": row, "key": code, "original": original_values(heads, values), "attrs": {"code": code, "name": text(values[1]), "document": document, "origin": "historico_importado"}, "problems": problems})
    issues["documents_shared_by_codes"] = sum(len(codes) > 1 for codes in documents.values())
    seen = set()
    for record in records:
        if record["key"] in seen:
            record["problems"].append("duplicate_supplier_code_preserved")
            issues["duplicate_supplier_code_preserved"] += 1
        seen.add(record["key"])
        if len(documents[record["attrs"]["document"]]) > 1:
            record["problems"].append("supplier_document_shared_by_codes")
    return records, {"unique_suppliers": len({r["key"] for r in records})}, dict(issues)


def read_workers(path):
    book = load_workbook(path, read_only=True, data_only=True)
    records, issues, seen = [], Counter(), set()
    try:
        sheet = book.active
        for row, values in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3):
            # Only the explicit lookup catalog in N:O; never execute VLOOKUP or RH formulas.
            if len(values) < 15 or not isinstance(values[13], (int, float)) or not values[14]:
                continue
            key = text(values[13])
            problems = []
            if key in seen:
                issues["duplicate_worker_registration"] += 1
                problems.append("duplicate_worker_registration")
            seen.add(key)
            records.append({"sheet": sheet.title, "row": row, "key": key, "original": {"registration": json_value(values[13]), "name": json_value(values[14])}, "attrs": {"registration": key, "name": text(values[14]), "origin": "historico_importado"}, "problems": problems})
    finally:
        book.close()
    if not records:
        raise PrivateDataError("Catálogo direto de matrículas não encontrado no boletim.")
    return records, {"unique_workers": len(seen), "bulletins_imported": 0}, dict(issues)


def read_movements(path):
    headers = ["Pedido Compra", "Data Lançamento", "Data do Documento", "Cod PN", "Nome", "Cod Item", "Desc Item", "Qtd", "Peso", "Deposito", "Nº Recebimento", "Data Recebimento", "Nota Fiscal de Entrada", "Chave de Acesso"]
    records, issues, orders, receipts, dates, weights = [], Counter(), set(), set(), [], defaultdict(set)
    for sheet, row, heads, values in workbook_rows(path, headers):
        problems = []
        order, receipt = text(values[0]), text(values[10])
        received = date_value(values[11], problems, "received_on")
        quantity = decimal_value(values[7], problems, "quantity")
        weight = decimal_value(values[8], problems, "weight")
        if quantity is not None and quantity < 0:
            problems.append("negative_quantity_preserved")
        key = text(values[13])
        if not key:
            problems.append("missing_invoice_key")
        elif len(key) != 44 or not key.isdigit():
            problems.append("invalid_invoice_key")
        if not order:
            problems.append("missing_purchase_order")
        if not receipt:
            problems.append("missing_receipt_number")
        attrs = {"source_sheet": sheet, "source_row": row, "purchase_order": order, "order_date": date_value(values[1], problems, "order_date"), "document_date": date_value(values[2], problems, "document_date"), "supplier_code": text(values[3]), "product_code": text(values[5]), "quantity": quantity, "reported_order_weight": weight, "depot": text(values[9]), "receipt_number": receipt, "received_on": received, "invoice_number": text(values[12]), "invoice_key": key, "original": original_values(heads, values), "problems": problems}
        records.append(attrs)
        issues.update(problems)
        if order:
            orders.add(order)
            if weight is not None:
                weights[order].add(weight)
        if receipt:
            receipts.add(receipt)
        if received:
            dates.append(received)
    summary = {"document_rows": len(records), "unique_orders": len(orders), "unique_document_receipts": len(receipts), "first_date": min(dates).isoformat() if dates else None, "last_date": max(dates).isoformat() if dates else None, "orders_with_constant_reported_weight": sum(len(v) == 1 for v in weights.values()), "trucks_identified": False, "measured_timestamps_available": False, "quantity_semantics_confirmed": False, "weight_semantics_confirmed": False}
    return records, summary, dict(issues)


def read_labor_days(path):
    records, issues, seen = [], Counter(), set()
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames != ["data", "dia_semana", "chapas_presentes", "chapas_operacao_cafe", "valor_pago_dia"]:
            raise PrivateDataError("Colunas do CSV incompatíveis com a versão do leitor.")
        for row, values in enumerate(reader, 2):
            problems = []
            day = date_value(values["data"], problems, "day")
            payroll = decimal_value(values["valor_pago_dia"], problems, "payroll_paid")
            try:
                count, coffee = int(values["chapas_presentes"]), int(values["chapas_operacao_cafe"])
            except (ValueError, TypeError):
                raise PrivateDataError("CSV contém contagem inválida; nenhuma linha do lote foi gravada.") from None
            if day is None or payroll is None or count < 0 or coffee < 0:
                raise PrivateDataError("CSV contém data ou valor inválido; nenhuma linha do lote foi gravada.")
            if day in seen:
                problems.append("duplicate_day_preserved")
            if coffee > count:
                problems.append("coffee_workers_above_total")
            seen.add(day)
            issues.update(problems)
            records.append({"source_row": row, "day": day, "weekday": values["dia_semana"], "worker_count": count, "coffee_worker_count": coffee, "payroll_paid": payroll, "original": values, "problems": problems})
    months = sorted({day.strftime("%Y-%m") for day in seen})
    summary = {"observed_dates": len(seen), "observed_months": months, "first_date": min(seen).isoformat() if seen else None, "last_date": max(seen).isoformat() if seen else None, "missing_months_documented": ["2025-08", "2025-12"], "half_days_available": False, "warehouse_allocation_available": False, "bulletin_cost_available": False}
    return records, summary, dict(issues)


READERS = {"products": read_products, "suppliers": read_suppliers, "workers": read_workers, "movements": read_movements, "labor_days": read_labor_days}


def record_outcomes(records):
    pending = sum(bool(record.get("problems")) for record in records)
    return {
        "accepted_rows": len(records) - pending,
        "pending_rows": pending,
        "rejected_rows": 0,
        "preserved_rows": len(records),
    }


def upsert_catalog(kind, records):
    model, key_name = {"products": (Product, "code"), "suppliers": (Supplier, "code"), "workers": (Worker, "registration")}[kind]
    first = {}
    for record in records:
        if record["attrs"][key_name]:
            first.setdefault(record["attrs"][key_name], record["attrs"])
    objects = [model(**attrs) for attrs in first.values()]
    if objects:
        fields = [f for f in first[next(iter(first))] if f != key_name]
        model.objects.bulk_create(objects, batch_size=1000, update_conflicts=True, update_fields=fields, unique_fields=[key_name])
    if kind == "products":
        products = dict(Product.objects.filter(code__in=first).values_list("code", "id"))
        pairs = {(r["attrs"]["code"], r["depot"]) for r in records if r["attrs"]["code"] and r["depot"]}
        ProductDeposit.objects.bulk_create([ProductDeposit(product_id=products[code], depot=depot) for code, depot in pairs], batch_size=1000, ignore_conflicts=True)


def import_source(kind, path, *, dry_run=False):
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    records, summary, issues = READERS[kind](path)
    if dry_run:
        return {"kind": kind, "status": "dry_run", "row_count": len(records), "summary": {**summary, **record_outcomes(records), "catalog_links_checked": False}, "issues": issues}
    with transaction.atomic():
        # Every process importing a kind takes the same PostgreSQL transaction lock.
        lock = int.from_bytes(hashlib.sha256(kind.encode()).digest()[:4], "big", signed=True)
        with connection.cursor() as cursor:
            cursor.execute("SELECT pg_advisory_xact_lock(%s)", [lock])
        existing = ImportBatch.objects.filter(kind=kind, file_hash=digest, importer_version=IMPORTER_VERSION).first()
        if existing and existing.active:
            return {"kind": kind, "status": "unchanged", "row_count": existing.row_count, "summary": existing.summary, "issues": existing.issues}
        ImportBatch.objects.filter(kind=kind, active=True).update(active=False)
        if kind in {"products", "suppliers", "workers"}:
            upsert_catalog(kind, records)
        if existing:
            existing.active = True
            existing.save(update_fields=["active"])
            return {"kind": kind, "status": "reactivated", "row_count": existing.row_count, "summary": existing.summary, "issues": existing.issues}
        batch = ImportBatch.objects.create(kind=kind, file_hash=digest, importer_version=IMPORTER_VERSION, source_name=path.name, row_count=len(records), summary=summary, issues=issues)
        if kind in {"products", "suppliers", "workers"}:
            SourceRow.objects.bulk_create([SourceRow(batch=batch, source_sheet=r["sheet"], source_row=r["row"], natural_key=r["key"], original=r["original"], problems=r["problems"]) for r in records], batch_size=1000)
        elif kind == "movements":
            active_catalog = ImportBatch.objects.filter(kind="products", active=True).first()
            # Current pairs come from the active source, not stale normalized associations.
            pairs = set(active_catalog.source_rows.values_list("natural_key", flat=True)) if active_catalog else set()
            current_codes = {pair.split("|", 1)[0] for pair in pairs}
            products = dict(Product.objects.filter(code__in=current_codes).values_list("code", "id"))
            active_suppliers = ImportBatch.objects.filter(kind="suppliers", active=True).first()
            supplier_codes = set(active_suppliers.source_rows.values_list("natural_key", flat=True)) if active_suppliers else set()
            suppliers = dict(Supplier.objects.filter(code__in=supplier_codes).values_list("code", "id"))
            counts = Counter(issues)
            for r in records:
                r["product_id"] = products.get(r["product_code"])
                r["supplier_id"] = suppliers.get(r["supplier_code"])
                if r["product_id"] is None:
                    r["problems"].append("product_missing_from_current_catalog")
                    counts["product_missing_from_current_catalog"] += 1
                elif f'{r["product_code"]}|{r["depot"]}' not in pairs:
                    r["problems"].append("product_depot_pair_missing_from_current_catalog")
                    counts["product_depot_pair_missing_from_current_catalog"] += 1
                if r["supplier_id"] is None:
                    r["problems"].append("supplier_missing_from_current_catalog")
                    counts["supplier_missing_from_current_catalog"] += 1
            HistoricalMovement.objects.bulk_create([HistoricalMovement(batch=batch, **r) for r in records], batch_size=1000)
            issues = dict(counts)
            batch.issues = issues
            batch.save(update_fields=["issues"])
        elif kind == "labor_days":
            HistoricalLaborDay.objects.bulk_create([HistoricalLaborDay(batch=batch, **r) for r in records], batch_size=1000)
        summary = {**summary, **record_outcomes(records), "catalog_links_checked": kind == "movements"}
        batch.summary = summary
        batch.save(update_fields=["summary"])
    return {"kind": kind, "status": "imported", "row_count": len(records), "summary": summary, "issues": issues}


def import_directory(directory, *, dry_run=False, kinds=None):
    directory = Path(directory).expanduser().resolve()
    if not directory.is_dir():
        raise PrivateDataError("Pasta privada de dados não encontrada. Configure PRIVATE_DATA_DIR ou --path.")
    results = []
    found = 0
    for kind in (kinds or SOURCE_FILES):
        path = directory / SOURCE_FILES[kind]
        if not path.is_file():
            results.append({"kind": kind, "status": "source_missing", "row_count": 0})
            continue
        found += 1
        try:
            results.append(import_source(kind, path, dry_run=dry_run))
        except (OSError, ValueError, KeyError, StopIteration) as error:
            if isinstance(error, PrivateDataError):
                raise
            raise PrivateDataError(f"Falha de leitura no conjunto {kind}; confira o arquivo privado.") from None
    if not found:
        raise PrivateDataError("Nenhuma fonte reconhecida foi encontrada na pasta privada.")
    return results
